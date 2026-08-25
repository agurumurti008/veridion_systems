"""
fcov_forge/generators/excel_template.py
=========================================
Generates a pre-formatted Excel workbook template for FCovForge.
Teams can fill this in instead of writing YAML by hand.

Also exports an existing FeatureModel to Excel format.
"""

from __future__ import annotations
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import (
        Font, PatternFill, Alignment, Border, Side, numbers
    )
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from core.model import FeatureModel, BinType


# ---------------------------------------------------------------------------
# Style constants
# ---------------------------------------------------------------------------

HEADER_FILL = "1E293B"
HEADER_FONT_COLOR = "E2E8F0"
SUBHEADER_FILL = "334155"
SUBHEADER_FONT_COLOR = "94A3B8"
STRIPE1 = "0F172A"
STRIPE2 = "1E293B"
ACCENT = "6366F1"
GREEN = "10B981"
RED = "EF4444"
YELLOW = "F59E0B"

BIN_TYPE_COLORS = {
    "values":     "1D4ED8",
    "range":      "7C3AED",
    "transition": "9D174D",
    "default":    "374151",
    "ignore":     "B45309",
    "illegal":    "991B1B",
    "wildcard":   "065F46",
    "auto":       "0E7490",
}


def _hfill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(bold=False, color="E2E8F0", size=10) -> Font:
    return Font(bold=bold, color=color, size=size, name="Calibri")


def _border() -> Border:
    side = Side(style="thin", color="334155")
    return Border(left=side, right=side, top=side, bottom=side)


def _header_row(ws, row: int, cols: list[tuple[str, int]], fill: str = HEADER_FILL):
    """Write a header row with labels and column widths."""
    for i, (label, width) in enumerate(cols, start=1):
        cell = ws.cell(row=row, column=i, value=label)
        cell.font = _font(bold=True, color=HEADER_FONT_COLOR)
        cell.fill = _hfill(fill)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = _border()
        ws.column_dimensions[get_column_letter(i)].width = width


def _data_row(ws, row: int, values: list, stripe: bool = False):
    fill_color = STRIPE1 if stripe else STRIPE2
    for i, val in enumerate(values, start=1):
        cell = ws.cell(row=row, column=i, value=val)
        cell.font = _font(color="CBD5E1")
        cell.fill = _hfill(fill_color)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = _border()


# ---------------------------------------------------------------------------
# Sheet builders
# ---------------------------------------------------------------------------

def _build_features_sheet(ws, model: FeatureModel):
    ws.title = "Features"
    ws.sheet_view.showGridLines = False
    cols = [
        ("feature_name", 25), ("description", 50),
        ("tags", 20), ("source_doc", 30),
    ]
    _header_row(ws, 1, cols)
    ws.row_dimensions[1].height = 20

    for i, feat in enumerate(model.features):
        _data_row(ws, i + 2, [
            feat.name,
            feat.description or "",
            ", ".join(feat.tags),
            feat.source_doc or "",
        ], stripe=i % 2 == 0)


def _build_covergroups_sheet(ws, model: FeatureModel):
    ws.title = "CoverGroups"
    ws.sheet_view.showGridLines = False
    cols = [
        ("feature_name", 22), ("cg_name", 25), ("clock", 18),
        ("condition", 18), ("per_instance", 12), ("auto_bin_max", 12),
        ("goal", 8), ("comment", 40), ("args", 30),
    ]
    _header_row(ws, 1, cols)

    row = 2
    for feat in model.features:
        for cg in feat.covergroups:
            args_str = ", ".join(f"{t} {n}" for t, n in cg.args)
            _data_row(ws, row, [
                feat.name, cg.name,
                cg.clock_event or "",
                cg.sample_condition or "",
                "TRUE" if cg.per_instance else "FALSE",
                cg.auto_bin_max or "",
                cg.goal or 100,
                cg.comment or "",
                args_str,
            ], stripe=row % 2 == 0)
            row += 1


def _build_coverpoints_sheet(ws, model: FeatureModel):
    ws.title = "CoverPoints"
    ws.sheet_view.showGridLines = False
    cols = [
        ("feature_name", 22), ("cg_name", 22), ("cp_name", 22),
        ("variable", 30), ("condition", 20), ("comment", 40),
    ]
    _header_row(ws, 1, cols)

    row = 2
    for feat in model.features:
        for cg in feat.covergroups:
            for cp in cg.coverpoints:
                _data_row(ws, row, [
                    feat.name, cg.name, cp.name,
                    cp.variable,
                    cp.condition or "",
                    cp.comment or "",
                ], stripe=row % 2 == 0)
                row += 1


def _build_bins_sheet(ws, model: FeatureModel):
    ws.title = "Bins"
    ws.sheet_view.showGridLines = False
    cols = [
        ("feature_name", 18), ("cg_name", 18), ("cp_name", 18),
        ("bin_name", 22), ("type", 12),
        ("values", 25), ("ranges", 25), ("transitions", 35),
        ("pattern", 20), ("auto_bin_max", 12),
        ("min_hits", 10), ("weight", 10), ("comment", 30),
    ]
    _header_row(ws, 1, cols)

    row = 2
    for feat in model.features:
        for cg in feat.covergroups:
            for cp in cg.coverpoints:
                for b in cp.bins:
                    # Serialize values
                    vals_str = ", ".join(str(v) for v in b.values) if b.values else ""
                    ranges_str = ", ".join(f"{lo}:{hi}" for lo, hi in b.ranges) if b.ranges else ""
                    trans_str = " | ".join(
                        " => ".join(str(s) for s in seq)
                        for seq in b.transitions
                    ) if b.transitions else ""

                    type_str = b.bin_type.value
                    cell_vals = [
                        feat.name, cg.name, cp.name, b.name, type_str,
                        vals_str, ranges_str, trans_str,
                        b.wildcard_pattern or "",
                        b.auto_bin_max or "",
                        b.min_hits if b.min_hits > 1 else "",
                        b.weight or "",
                        b.comment or "",
                    ]
                    _data_row(ws, row, cell_vals, stripe=row % 2 == 0)
                    # Color type cell
                    type_cell = ws.cell(row=row, column=5)
                    color = BIN_TYPE_COLORS.get(type_str, "374151")
                    type_cell.fill = _hfill(color)
                    type_cell.font = _font(bold=True, color="FFFFFF")
                    row += 1


def _build_cg_crosses_sheet(ws, model: FeatureModel):
    ws.title = "CG_Crosses"
    ws.sheet_view.showGridLines = False
    cols = [
        ("feature_name", 22), ("cg_name", 22), ("cross_name", 25),
        ("coverpoints", 40), ("exclude_bins", 40), ("comment", 40),
    ]
    _header_row(ws, 1, cols)

    row = 2
    for feat in model.features:
        for cg in feat.covergroups:
            for cross in cg.crosses:
                _data_row(ws, row, [
                    feat.name, cg.name, cross.name,
                    ", ".join(cross.coverpoints),
                    ", ".join(cross.exclude_bins),
                    cross.comment or "",
                ], stripe=row % 2 == 0)
                row += 1


def _build_crosses_sheet(ws, model: FeatureModel):
    ws.title = "Crosses"
    ws.sheet_view.showGridLines = False
    cols = [
        ("cross_name", 30), ("targets", 60),
        ("exclude_combinations", 50), ("comment", 40), ("goal", 8),
    ]
    _header_row(ws, 1, cols)

    row = 2
    for fc in model.feature_crosses:
        targets_str = " | ".join(t.address for t in fc.targets)
        _data_row(ws, row, [
            fc.name, targets_str, "", fc.comment or "", fc.goal or 100,
        ], stripe=row % 2 == 0)
        row += 1


def _build_instructions_sheet(ws):
    ws.title = "Instructions"
    ws.sheet_view.showGridLines = False

    instructions = [
        ("FCovForge Excel Format — Quick Reference", True, ACCENT),
        ("", False, STRIPE1),
        ("SHEETS OVERVIEW", True, SUBHEADER_FILL),
        ("Features      — One row per feature (top-level)", False, STRIPE2),
        ("CoverGroups   — One row per covergroup", False, STRIPE1),
        ("CoverPoints   — One row per coverpoint", False, STRIPE2),
        ("Bins          — One row per bin (all types)", False, STRIPE1),
        ("CG_Crosses    — Crosses WITHIN a covergroup", False, STRIPE2),
        ("Crosses       — Feature-level crosses (FCovForge innovation)", False, STRIPE1),
        ("", False, STRIPE1),
        ("BIN TYPES", True, SUBHEADER_FILL),
        ("values     — bins b = {1, 2, 3}          (put values in 'values' column, comma-separated)", False, STRIPE2),
        ("range      — bins b = {[0:15]}            (put lo:hi in 'ranges' column)", False, STRIPE1),
        ("transition — bins b = (A => B => C)      (put A => B => C in 'transitions'; | separates seqs)", False, STRIPE2),
        ("wildcard   — wildcard bins b = {4'b1?0?} (put pattern in 'pattern' column)", False, STRIPE1),
        ("auto       — auto bins                   (set auto_bin_max)", False, STRIPE2),
        ("default    — bins b = default            (no values needed)", False, STRIPE1),
        ("ignore     — ignore_bins b = {255}       (values/ranges as usual)", False, STRIPE2),
        ("illegal    — illegal_bins b = {[256:511]}(values/ranges as usual)", False, STRIPE1),
        ("", False, STRIPE1),
        ("FEATURE CROSSES (Crosses sheet)", True, SUBHEADER_FILL),
        ("Targets column: pipe-separated addresses", False, STRIPE2),
        ("  CP level:  feat1.cg1.cp1 | feat2.cg2.cp2", False, STRIPE1),
        ("  CG level:  feat1.cg1 | feat2.cg2", False, STRIPE2),
        ("  Feat level: feat1 | feat2", False, STRIPE1),
        ("  3-way:     feat1.cg1.cp1 | feat2.cg2.cp2 | feat3.cg3.cp3", False, STRIPE2),
    ]

    for row, (text, bold, fill_color) in enumerate(instructions, start=1):
        cell = ws.cell(row=row, column=1, value=text)
        cell.font = _font(bold=bold, color="E2E8F0", size=11 if bold else 10)
        cell.fill = _hfill(fill_color)
        cell.alignment = Alignment(vertical="center")
        ws.row_dimensions[row].height = 18

    ws.column_dimensions["A"].width = 90


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def export_to_excel(model: FeatureModel, output: Path) -> None:
    """Export a FeatureModel to the FCovForge Excel format."""
    if not OPENPYXL_AVAILABLE:
        raise ImportError("openpyxl required: pip install openpyxl")

    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    _build_instructions_sheet(wb.create_sheet())
    _build_features_sheet(wb.create_sheet(), model)
    _build_covergroups_sheet(wb.create_sheet(), model)
    _build_coverpoints_sheet(wb.create_sheet(), model)
    _build_bins_sheet(wb.create_sheet(), model)
    _build_cg_crosses_sheet(wb.create_sheet(), model)
    _build_crosses_sheet(wb.create_sheet(), model)

    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    print(f"[Excel Export] Written: {output}")


def generate_blank_template(output: Path) -> None:
    """Generate a blank Excel template with example rows."""
    from core.model import (
        FeatureModel, Feature, CoverGroup, CoverPoint, CoverBin,
        FeatureCross, FeatureCrossTarget
    )

    # Build a minimal example model
    model = FeatureModel(
        project_name="my_project",
        features=[
            Feature(
                name="example_feature_1",
                description="Replace with your feature description",
                tags=["example", "uart"],
                covergroups=[
                    CoverGroup(
                        name="example_cg_1",
                        clock_event="posedge clk",
                        comment="Example covergroup",
                        coverpoints=[
                            CoverPoint(
                                name="example_cp_1",
                                variable="your_signal_here",
                                bins=[
                                    CoverBin(name="BIN_ZERO", bin_type=BinType.VALUES, values=[0]),
                                    CoverBin(name="BIN_ONE", bin_type=BinType.VALUES, values=[1]),
                                    CoverBin(name="BIN_RANGE", bin_type=BinType.RANGE, ranges=[(2, 15)]),
                                ],
                            )
                        ],
                    )
                ],
            ),
            Feature(
                name="example_feature_2",
                description="Replace with your second feature",
                covergroups=[
                    CoverGroup(
                        name="example_cg_2",
                        coverpoints=[
                            CoverPoint(
                                name="state_cp",
                                variable="state_machine",
                                bins=[
                                    CoverBin(
                                        name="BIN_IDLE_TO_ACTIVE",
                                        bin_type=BinType.TRANSITION,
                                        transitions=[["IDLE", "ACTIVE"]],
                                    ),
                                ],
                            )
                        ],
                    )
                ],
            ),
        ],
        feature_crosses=[
            FeatureCross(
                name="example_cross",
                targets=[
                    FeatureCrossTarget("example_feature_1.example_cg_1.example_cp_1"),
                    FeatureCrossTarget("example_feature_2.example_cg_2.state_cp"),
                ],
                comment="Example feature-level cross",
                goal=100,
            )
        ],
    )

    export_to_excel(model, output)
